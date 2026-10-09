"""FAT/SNF workflow integration tests (real MySQL)."""
from datetime import date, timedelta
from decimal import Decimal as D

import pytest

from conftest import (enable_fat_snf, login_owner, login_staff, make_owner, make_staff,
                      make_vendor, now_date_slot, q, reference_chart_payload)


@pytest.fixture()
def dairy(client):
    owner = make_owner()
    make_vendor(owner, 101, "प्रदीप जगदाळे", "cow")
    make_vendor(owner, 102, "सुनील पाटील", "buffalo")
    make_vendor(owner, 103, "राम शिंदे", "both")
    login_owner(client, owner)
    r = client.post("/api/fat_snf/charts", json=reference_chart_payload("cow"))
    assert r.status_code == 201, r.get_json()
    return {"owner": owner, "chart_id": r.get_json()["chart_id"]}


def entry(client, vendor=101, fat="3.5", snf="8.5", litres="18", milk_type="cow", **extra):
    d, s = now_date_slot()
    body = {"vendor_id": vendor, "milk_type": milk_type, "fat": fat, "snf": snf,
            "litres": litres, "date": d, "slot": s}
    body.update(extra)
    return client.post("/api/fat_snf/entries", json=body)


# ----------------------------------------------------------------------
# calculation mode
# ----------------------------------------------------------------------
class TestMode:
    def test_default_off_and_owner_toggle_is_audited(self, client, dairy):
        assert client.get("/api/fat_snf/calculation_mode").get_json()["calculation_mode"] == "general"
        r = client.post("/api/fat_snf/calculation_mode", json={"calculation_mode": "fat_snf"})
        assert r.status_code == 200 and r.get_json()["calculation_mode"] == "fat_snf"
        audit = q("SELECT * FROM dm_audit_log WHERE action='calculation_mode_change'", one=True)
        assert audit and '"general"' in audit["old_values"] and '"fat_snf"' in audit["new_values"]

    def test_staff_cannot_change_mode(self, client, dairy):
        staff = make_staff(dairy["owner"])
        login_staff(client, dairy["owner"], staff)
        r = client.post("/api/fat_snf/calculation_mode", json={"calculation_mode": "fat_snf"})
        assert r.status_code == 403
        assert client.get("/api/fat_snf/calculation_mode").get_json()["calculation_mode"] == "general"

    def test_invalid_mode(self, client, dairy):
        assert client.post("/api/fat_snf/calculation_mode",
                           json={"calculation_mode": "x"}).status_code == 400

    def test_saving_requires_mode_on(self, client, dairy):
        r = entry(client)
        assert r.status_code == 409 and r.get_json()["code"] == "mode_off"

    def test_legacy_collection_blocked_when_on(self, client, dairy):
        enable_fat_snf(dairy["owner"])
        r = client.post("/submit_milk_ajax", json={"vendor_id": 101, "milk_type": "cow", "quantity": 5})
        assert r.status_code == 409 and r.get_json()["code"] == "mode_on"
        assert q("SELECT COUNT(*) n FROM milk_collection", one=True)["n"] == 0
        # legacy pages redirect to FAT/SNF equivalents, ?legacy=1 opens the old one
        assert client.get("/milk_collection").headers["Location"].endswith("/fat_snf/milk_collection")
        assert client.get("/payment?legacy=1").status_code == 200

    def test_legacy_collection_still_works_when_off(self, client, dairy):
        d, s = now_date_slot()
        r = client.post("/submit_milk_ajax", json={"vendor_id": 101, "milk_type": "cow",
                                                   "quantity": 5, "date": d, "slot": s})
        assert r.status_code == 200, r.get_json()
        assert q("SELECT COUNT(*) n FROM milk_collection", one=True)["n"] == 1


# ----------------------------------------------------------------------
# rate charts
# ----------------------------------------------------------------------
class TestCharts:
    def test_chart_stored_with_settings(self, client, dairy):
        assert q("SELECT COUNT(*) n FROM fat_snf_chart_values WHERE chart_id=%s",
                 (dairy["chart_id"],), one=True)["n"] == 95
        c = q("SELECT * FROM fat_snf_charts WHERE id=%s", (dairy["chart_id"],), one=True)
        assert c["std_rate"] == D("27.00") and c["fat_rate_step"] == D("0.30") and c["version"] == 1

    def test_duplicate_active_chart_conflict(self, client, dairy):
        r = client.post("/api/fat_snf/charts", json=reference_chart_payload("cow"))
        assert r.status_code == 409

    def test_incomplete_matrix_rejected(self, client, dairy):
        p = reference_chart_payload("buffalo")
        p["matrix"] = p["matrix"][:-1]
        r = client.post("/api/fat_snf/charts", json=p)
        assert r.status_code == 400 and q(
            "SELECT COUNT(*) n FROM fat_snf_charts", one=True)["n"] == 1

    def test_preview_api(self, client, dairy):
        from conftest import REFERENCE_SETTINGS
        r = client.post("/api/fat_snf/charts/preview", json=REFERENCE_SETTINGS)
        body = r.get_json()
        assert r.status_code == 200 and len(body["matrix"]) == 95

    def test_edit_cell_keeps_historical_entry_rate(self, client, dairy):
        enable_fat_snf(dairy["owner"])
        assert entry(client).status_code == 201
        p = reference_chart_payload("cow")
        for c in p["matrix"]:
            if D(c["fat"]) == D("3.5") and D(c["snf"]) == D("8.5"):
                c["rate"] = "30.00"
        r = client.put(f"/api/fat_snf/charts/{dairy['chart_id']}", json=p)
        assert r.status_code == 200 and r.get_json()["changed_cells"] == 1
        e = q("SELECT rate, amount FROM fat_snf_milk_collection", one=True)
        assert e["rate"] == D("27.00") and e["amount"] == D("486.00")
        # the change is audited with old and new value
        a = q("SELECT * FROM dm_audit_log WHERE action='rate_chart_update'", one=True)
        assert "27.00" in a["old_values"] and "30.00" in a["new_values"]

    def test_revision_and_effective_dates(self, client, dairy):
        enable_fat_snf(dairy["owner"])
        today = date.today()
        p = reference_chart_payload("cow", effective_from=today.isoformat(), name="v2")
        for c in p["matrix"]:
            c["rate"] = str(D(c["rate"]) + 1)
        p["revised_from"] = dairy["chart_id"]
        r = client.post("/api/fat_snf/charts", json=p)
        assert r.status_code == 201
        assert q("SELECT version FROM fat_snf_charts WHERE id=%s", (r.get_json()["chart_id"],),
                 one=True)["version"] == 2
        y = (today - timedelta(days=1)).isoformat()
        old = client.get(f"/api/fat_snf/rate?vendor_id=101&milk_type=cow&fat=3.5&snf=8.5&date={y}")
        new = client.get(f"/api/fat_snf/rate?vendor_id=101&milk_type=cow&fat=3.5&snf=8.5&date={today}")
        assert old.get_json()["rate"] == "27.00" and new.get_json()["rate"] == "28.00"

    def test_deactivate(self, client, dairy):
        r = client.post(f"/api/fat_snf/charts/{dairy['chart_id']}/status", json={"active": False})
        assert r.status_code == 200
        r = client.get("/api/fat_snf/rate?vendor_id=101&milk_type=cow&fat=3.5&snf=8.5")
        assert r.status_code == 422 and r.get_json()["code"] == "no_chart"

    def test_pages_render(self, client, dairy):
        assert client.get("/fat_snf/rate-chart").status_code == 200
        assert client.get("/fat_snf/rate-chart/create").status_code == 200
        r = client.get(f"/fat_snf/rate-chart/{dairy['chart_id']}")
        assert r.status_code == 200 and "Reference" in r.get_data(as_text=True)
        assert client.get(f"/fat_snf/rate-chart/create?revise={dairy['chart_id']}").status_code == 200


# ----------------------------------------------------------------------
# collection
# ----------------------------------------------------------------------
class TestCollection:
    @pytest.fixture(autouse=True)
    def _on(self, dairy):
        enable_fat_snf(dairy["owner"])

    def test_save_calculates_rate_and_amount(self, client, dairy):
        r = entry(client, fat="3.5", snf="8.5", litres="18")
        e = r.get_json()["entry"]
        assert r.status_code == 201
        assert e["rate"] == "27.00" and e["amount"] == "486.00" and e["status"] == "saved"
        r = entry(client, vendor=103, fat="3.3", snf="8.4", litres="12")
        assert r.get_json()["entry"]["amount"] == "313.20"   # 12 x 26.10

    def test_rate_and_amount_never_taken_from_client(self, client, dairy):
        r = entry(client, rate="999", amount="99999")
        assert r.get_json()["entry"]["rate"] == "27.00"

    def test_duplicate_business_key(self, client, dairy):
        assert entry(client).status_code == 201
        r = entry(client, litres="5")
        assert r.status_code == 409 and r.get_json()["code"] == "duplicate"
        assert q("SELECT COUNT(*) n FROM fat_snf_milk_collection", one=True)["n"] == 1

    def test_idempotent_retry(self, client, dairy):
        first = entry(client, request_id="req-123")
        again = entry(client, request_id="req-123")
        assert first.status_code == 201 and again.status_code == 200
        assert again.get_json()["replayed"] is True
        assert q("SELECT COUNT(*) n FROM fat_snf_milk_collection", one=True)["n"] == 1
        st = client.get("/api/fat_snf/requests/req-123").get_json()
        assert st["saved"] is True

    def test_missing_combination(self, client, dairy):
        r = entry(client, fat="6.0")
        assert r.status_code == 422 and r.get_json()["code"] == "no_cell"
        assert "दर उपलब्ध नाही" in r.get_json()["message"]

    def test_no_chart_for_buffalo(self, client, dairy):
        r = entry(client, vendor=102, milk_type="buffalo")
        assert r.status_code == 422 and r.get_json()["code"] == "no_chart"

    def test_milk_type_must_match_vendor(self, client, dairy):
        r = entry(client, vendor=101, milk_type="buffalo")
        assert r.status_code == 400 and r.get_json()["field"] == "milk_type"

    @pytest.mark.parametrize("field,value", [("fat", "abc"), ("snf", "-1"), ("litres", "0"),
                                             ("litres", "-3"), ("fat", "3.555"), ("vendor", 0)])
    def test_validation(self, client, dairy, field, value):
        if field == "vendor":
            r = entry(client, vendor=value)
        else:
            r = entry(client, **{field: value})
        assert r.status_code == 400

    def test_unknown_vendor(self, client, dairy):
        assert entry(client, vendor=999).status_code == 404

    def test_future_date_and_wrong_slot(self, client, dairy):
        d, s = now_date_slot()
        tomorrow = (date.fromisoformat(d) + timedelta(days=1)).isoformat()
        assert entry(client, date=tomorrow).status_code == 400
        other = "evening" if s == "morning" else "morning"
        r = entry(client, slot=other)
        assert r.status_code == 409 and r.get_json()["code"] == "slot_changed"

    def test_backdate_owner_only(self, client, dairy):
        d, _ = now_date_slot()
        y = (date.fromisoformat(d) - timedelta(days=3)).isoformat()
        r = entry(client, date=y, slot="morning", backdated=True)
        assert r.status_code == 201
        assert q("SELECT action FROM dm_audit_log WHERE action='fsn_entry_create'", one=True)
        staff = make_staff(dairy["owner"])
        login_staff(client, dairy["owner"], staff)
        r = entry(client, vendor=103, date=y, slot="morning", backdated=True)
        assert r.status_code == 403

    def test_warning_needs_confirmation(self, client, dairy):
        r = entry(client, litres="150")
        assert r.status_code == 409 and r.get_json()["code"] == "needs_confirmation"
        assert q("SELECT COUNT(*) n FROM fat_snf_milk_collection", one=True)["n"] == 0
        r = entry(client, litres="150", confirm_warnings=True)
        assert r.status_code == 201

    def test_staff_can_collect_and_staff_id_stored(self, client, dairy):
        staff = make_staff(dairy["owner"])
        login_staff(client, dairy["owner"], staff)
        assert client.get("/fat_snf/milk_collection").status_code == 200
        assert entry(client).status_code == 201
        assert q("SELECT staff_id FROM fat_snf_milk_collection", one=True)["staff_id"] == staff

    def test_finalize_and_correction_rules(self, client, dairy):
        d, s = now_date_slot()
        eid = entry(client).get_json()["entry"]["id"]
        r = client.post("/api/fat_snf/finalize", json={"date": d, "slot": s})
        assert r.get_json()["finalized"] == 1
        assert client.post("/api/fat_snf/finalize", json={"date": d, "slot": s}).get_json()["finalized"] == 0
        # staff cannot change a final entry
        staff = make_staff(dairy["owner"])
        login_staff(client, dairy["owner"], staff)
        assert client.put(f"/api/fat_snf/entries/{eid}",
                          json={"fat": "3.6", "snf": "8.5", "litres": "18"}).status_code == 403
        # owner needs a reason, then the rate is re-read from the chart
        login_owner(client, dairy["owner"])
        assert client.put(f"/api/fat_snf/entries/{eid}",
                          json={"fat": "3.6", "snf": "8.5", "litres": "18"}).status_code == 400
        r = client.put(f"/api/fat_snf/entries/{eid}",
                       json={"fat": "3.6", "snf": "8.5", "litres": "18", "reason": "चूक"})
        body = r.get_json()
        assert r.status_code == 200 and body["entry"]["rate"] == "27.30"
        assert body["old_amount"] == "486.00" and body["entry"]["amount"] == "491.40"

    def test_stale_edit_detected(self, client, dairy):
        e = entry(client).get_json()["entry"]
        r = client.put(f"/api/fat_snf/entries/{e['id']}",
                       json={"fat": "3.6", "snf": "8.5", "litres": "18",
                             "updated_at": "2000-01-01T00:00:00"})
        assert r.status_code == 409 and r.get_json()["code"] == "stale"

    def test_delete_requires_owner_and_reason(self, client, dairy):
        eid = entry(client).get_json()["entry"]["id"]
        assert client.delete(f"/api/fat_snf/entries/{eid}", json={}).status_code == 400
        staff = make_staff(dairy["owner"])
        login_staff(client, dairy["owner"], staff)
        assert client.delete(f"/api/fat_snf/entries/{eid}", json={"reason": "x"}).status_code == 403
        login_owner(client, dairy["owner"])
        assert client.delete(f"/api/fat_snf/entries/{eid}", json={"reason": "x"}).status_code == 200
        assert q("SELECT action FROM dm_audit_log WHERE action='fsn_entry_delete'", one=True)

    def test_list_and_pages(self, client, dairy):
        entry(client)
        d, s = now_date_slot()
        body = client.get(f"/api/fat_snf/entries?date={d}&slot={s}").get_json()
        assert body["totals"]["count"] == 1 and body["totals"]["amount"] == "486.00"
        assert client.get("/fat_snf/milk_collection").status_code == 200
        assert client.get(f"/fat_snf/edit_entry?date={d}").status_code == 200
        assert client.get("/fat_snf/milk_summary").status_code == 200


# ----------------------------------------------------------------------
# calculation, receipts, settlement
# ----------------------------------------------------------------------
class TestSettlement:
    @pytest.fixture()
    def period(self, client, dairy):
        """Ten days of cow milk for vendor 101 plus food sack and advances."""
        owner = dairy["owner"]
        enable_fat_snf(owner)
        start = date(2026, 9, 21)
        for i in range(10):
            d = start + timedelta(days=i)
            for slot, fat, litres in (("morning", "3.5", "18"), ("evening", "3.4", "15")):
                r = entry(client, fat=fat, snf="8.5", litres=litres, date=d.isoformat(),
                          slot=slot, backdated=True)
                assert r.status_code == 201, r.get_json()
        rate_id = q("INSERT INTO food_sack_rates (user_id, name, rate, date_from) "
                    "VALUES (%s,'नेचर',1900,'2026-01-01')", (owner,))
        q("INSERT INTO food_sack (vendor_id, user_id, date, sack_qty, sack_rate_id, total_cost, "
          "sack_rate) VALUES (101,%s,'2026-09-25',2,%s,3800,1900)", (owner, rate_id))
        for t, amt, d in (("advance", 5000, "2026-08-01"), ("deduction", 1000, "2026-09-28"),
                          ("deduction", 500, "2026-08-15")):
            q("INSERT INTO advance_transactions (vendor_id, user_id, transaction_type, amount, "
              "transaction_date) VALUES (101,%s,%s,%s,%s)", (owner, t, amt, d))
        return owner

    def test_statement_math(self, client, period):
        from fat_snf import service as S
        with client.application.test_request_context():
            vendor = S.get_vendor(period, 101)
            st = S.build_statements(period, date(2026, 9, 21), date(2026, 9, 30), [vendor])[0]
        cow = st["sections"][0]
        assert [s["type"] for s in st["sections"]] == ["cow"]          # cow-only vendor
        assert cow["morning"]["litres"] == D("180.00") and cow["morning"]["amount"] == D("4860.00")
        assert cow["evening"]["amount"] == D("4005.00")               # 15 x 26.70 x 10
        assert cow["litres"] == D("330.00") and cow["amount"] == D("8865.00")
        assert cow["avg_fat"] == D("3.45")                            # (18*3.5+15*3.4)/33
        assert st["food"]["total"] == D("3800.00")
        assert st["period_deduction"] == D("1000.00")                 # only Sept deduction
        assert st["remaining_advance"] == D("3500.00")                # 5000-1000-500
        assert st["final"] == D("4065.00")                            # 8865-3800-1000
        assert len(cow["rows"]) == 10

    def test_receipt_page(self, client, period):
        r = client.post("/fat_snf/receipt", data={"from_date": "2026-09-21", "to_date": "2026-09-30",
                                                  "vendor_id": "101"})
        html = r.get_data(as_text=True)
        assert r.status_code == 200
        for text in ("प्रदीप जगदाळे", "गाय", "फॅट %", "SNF %", "दर / लि.", "486.00", "400.50",
                     "₹8,865.00", "नेचर (₹1900/गोणी) × 2 = ₹3,800.00", "₹4,065.00", "₹3,500.00",
                     "21/09/2026 ते 30/09/2026"):
            assert text in html, text
        assert "म्हैस</th>" not in html                               # no buffalo section

    def test_both_type_vendor_gets_both_sections(self, client, period):
        r = client.post("/fat_snf/receipt", data={"from_date": "2026-09-21", "to_date": "2026-09-30",
                                                  "vendor_id": "103"})
        html = r.get_data(as_text=True)
        assert html.count('class="type-head"') == 2

    def test_all_vendor_receipts(self, client, period):
        r = client.post("/fat_snf/receipt", data={"from_date": "2026-09-21", "to_date": "2026-09-30",
                                                  "vendor_id": "all"})
        assert r.get_data(as_text=True).count("<article") == 1   # only vendors with activity

    def test_receipt_rejects_bad_range(self, client, period):
        r = client.post("/fat_snf/receipt", data={"from_date": "2026-09-30", "to_date": "2026-09-01"})
        assert "सुरुवातीची तारीख" in r.get_data(as_text=True)

    def test_calculation_page_and_csv(self, client, period):
        r = client.get("/fat_snf/calculation?from_date=2026-09-21&to_date=2026-09-30")
        html = r.get_data(as_text=True)
        assert r.status_code == 200 and "₹4,065.00" in html
        r = client.get("/fat_snf/calculation?from_date=2026-09-21&to_date=2026-09-30&export=csv")
        assert r.mimetype == "text/csv" and "4065.00" in r.get_data(as_text=True)

    def test_legacy_pages_include_fat_snf_after_switch_back(self, client, period):
        q("UPDATE fat_snf_calculation_settings SET calculation_mode='general'")
        r = client.get("/payment?start_date=2026-09-21&end_date=2026-09-30")
        assert r.status_code == 200 and "₹4065" in r.get_data(as_text=True)

    def test_mixed_period_adds_legacy_litre_rate_milk(self, client, period):
        owner = period
        q("INSERT INTO milk_rates (user_id, animal, rate, date_from) VALUES (%s,'cow',30,'2026-01-01')",
          (owner,))
        q("INSERT INTO milk_collection (vendor_id, user_id, date, slot, milk_type, quantity) "
          "VALUES (101,%s,'2026-09-20','morning','cow',10)", (owner,))
        from fat_snf import service as S
        with client.application.test_request_context():
            st = S.build_statements(owner, date(2026, 9, 20), date(2026, 9, 30),
                                    [S.get_vendor(owner, 101)])[0]
        assert st["general_total"] == D("300.00")
        assert st["final"] == D("4365.00")

    def test_customer_sees_own_fat_snf_receipt_only(self, client, period):
        with client.session_transaction() as s:
            s.clear()
            s.update(loggedin=True, role="customer", id=period, owner_id=period, vendor_id=101,
                     customer_id=1, customer_name="प्रदीप", dairy_name="Test")
        r = client.post("/receipt/102", data={"from_date": "2026-09-21", "to_date": "2026-09-30"})
        html = r.get_data(as_text=True)
        assert r.status_code == 200 and "प्रदीप जगदाळे" in html and "सुनील पाटील" not in html

    def test_bank_report_generates(self, client, period):
        q("INSERT INTO bank_settings (user_id, account_holder_name, bank_name, account_no, "
          "ifsc_code) VALUES (%s,'Owner','SBI','1','SBIN0001')", (period,))
        r = client.post("/generate_bank_report", data={"from_date": "2026-09-21",
                                                       "to_date": "2026-09-30"})
        assert r.status_code == 200 and r.mimetype == "application/zip"
        import io
        import zipfile
        from openpyxl import load_workbook
        z = zipfile.ZipFile(io.BytesIO(r.data))
        ws = load_workbook(io.BytesIO(z.read("Other_Bank_Report.xlsx"))).active
        amounts = [row[5] for row in ws.iter_rows(min_row=6, values_only=True)]
        assert 4065 in amounts

    def test_per_vendor_legacy_receipt_no_longer_crashes(self, client, period):
        q("UPDATE fat_snf_calculation_settings SET calculation_mode='general'")
        r = client.post("/receipt/101", data={"from_date": "2026-09-21", "to_date": "2026-09-30"})
        assert r.status_code == 200
